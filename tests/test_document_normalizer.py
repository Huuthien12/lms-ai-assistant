from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from deeptutor_integration.contracts import SourceDocument
from deeptutor_integration.document_normalizer import DocumentNormalizer, source_sha256
from deeptutor_integration.errors import DeepTutorError


class DocumentNormalizerTests(unittest.TestCase):
    def source(self, filename="notes.md", mime_type="text/markdown", content=b"# Notes\r\n\r\nText"):
        return SourceDocument(
            document_id="resource-1", course_id="INT1339", filename=filename,
            mime_type=mime_type, source="moodle", content=content,
            metadata={"course_name": "Python", "token": "private", "download_url": "https://private"},
        )

    def test_normalizes_markdown_and_preserves_safe_provenance(self):
        result = DocumentNormalizer().normalize(self.source())
        self.assertEqual(result.markdown, "# Notes\n\nText")
        self.assertEqual(result.normalizer_version, "markdown-v1")
        self.assertEqual(result.document_id, "resource-1")
        self.assertEqual(result.course_id, "INT1339")
        self.assertEqual(result.original_filename, "notes.md")
        self.assertEqual(result.original_mime_type, "text/markdown")
        self.assertEqual(result.source, "moodle")
        self.assertEqual(result.metadata, {"course_name": "Python"})
        self.assertNotIn("private", str(result.model_dump()))

    def test_bom_is_removed(self):
        self.assertEqual(DocumentNormalizer().normalize(self.source(content=b"\xef\xbb\xbf# Notes")).markdown, "# Notes")

    def test_hash_is_deterministic_and_tracks_original_bytes(self):
        self.assertEqual(source_sha256(b"same"), source_sha256(b"same"))
        self.assertNotEqual(source_sha256(b"first"), source_sha256(b"second"))

    def test_empty_bytes_and_blank_markdown_are_rejected(self):
        with self.assertRaisesRegex(DeepTutorError, "non-empty"):
            source_sha256(b"")
        with self.assertRaisesRegex(DeepTutorError, "must not be blank"):
            DocumentNormalizer().normalize(self.source(content=b" \t\r\n"))

    def test_invalid_utf8_is_rejected_without_exposing_content(self):
        with self.assertRaises(DeepTutorError) as raised:
            DocumentNormalizer().normalize(self.source(content=b"secret-content-\xff"))
        self.assertEqual(raised.exception.code, "invalid_document")
        self.assertNotIn("secret-content", str(raised.exception))

    def test_unsupported_and_mismatched_formats_are_rejected(self):
        with self.assertRaisesRegex(DeepTutorError, "not supported"):
            DocumentNormalizer().normalize(self.source(filename="notes.exe", mime_type="application/octet-stream"))
        with self.assertRaisesRegex(DeepTutorError, "does not match"):
            DocumentNormalizer().normalize(self.source(filename="notes.md", mime_type="application/pdf"))

    def test_ppt_is_unsupported(self):
        with self.assertRaises(DeepTutorError) as raised:
            DocumentNormalizer().normalize(self.source(filename="notes.ppt", mime_type="application/vnd.ms-powerpoint", content=b"binary"))
        self.assertEqual((raised.exception.code, raised.exception.status_code), ("unsupported_document", 415))

    @staticmethod
    def pdf_bytes(text: str) -> bytes:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        output = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, item in enumerate(objects, 1):
            offsets.append(len(output))
            output.extend(f"{index} 0 obj\n".encode("ascii") + item + b"\nendobj\n")
        xref = len(output)
        output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
        output.extend(b"".join(f"{offset:010} 00000 n \n".encode("ascii") for offset in offsets[1:]))
        output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
        return bytes(output)

    @staticmethod
    def docx_bytes(text: str) -> bytes:
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr(
                "[Content_Types].xml",
                """<?xml version=\"1.0\"?><Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\"><Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/><Default Extension=\"xml\" ContentType=\"application/xml\"/><Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/></Types>""",
            )
            archive.writestr(
                "_rels/.rels",
                """<?xml version=\"1.0\"?><Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"><Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/></Relationships>""",
            )
            archive.writestr(
                "word/document.xml",
                f"""<?xml version=\"1.0\"?><w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>""",
            )
        return output.getvalue()

    @staticmethod
    def pptx_bytes(text: str) -> bytes:
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.add_textbox(0, 0, 9144000, 914400).text_frame.text = text
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    def test_real_markitdown_conversions_preserve_text_and_original_hash(self):
        cases = (
            ("notes.pdf", "application/pdf", self.pdf_bytes("DeepTutor PDF normalization test"), "markitdown-pdf-v1", "DeepTutor PDF normalization test"),
            ("notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", self.docx_bytes("DeepTutor DOCX normalization test"), "markitdown-docx-v1", "DeepTutor DOCX normalization test"),
            ("notes.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", self.pptx_bytes("DeepTutor PPTX normalization test"), "markitdown-pptx-v1", "DeepTutor PPTX normalization test"),
        )
        for filename, mime_type, content, version, expected_text in cases:
            with self.subTest(filename=filename):
                source = self.source(filename=filename, mime_type=mime_type, content=content)
                result = DocumentNormalizer().normalize(source)
                self.assertIn(expected_text, result.markdown)
                self.assertTrue(result.markdown.strip())
                self.assertEqual(result.sha256, source_sha256(content))
                self.assertEqual(result.normalizer_version, version)

    def test_converter_errors_are_sanitized_and_temp_files_are_removed(self):
        paths: list[Path] = []

        def fail(path: Path) -> str:
            paths.append(path)
            raise RuntimeError("secret-content C:/private/document.pdf")

        with patch("deeptutor_integration.document_normalizer._convert_markitdown", side_effect=fail):
            with self.assertRaises(DeepTutorError) as raised:
                DocumentNormalizer().normalize(self.source(filename="notes.pdf", mime_type="application/pdf", content=b"PDF"))
        self.assertEqual(raised.exception.code, "normalization_failed")
        self.assertNotIn("secret-content", str(raised.exception))
        self.assertFalse(paths[0].exists())


if __name__ == "__main__":
    unittest.main()
